"""Генератор презентации KnovaLand в формате .pptx.

Требует: pip install python-pptx
Запуск:   python tools/make_pptx.py
Результат: docs/KnovaLand_slides.pptx
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

BG = RGBColor(0x0B, 0x11, 0x20)
TITLE = RGBColor(0xA5, 0xB4, 0xFC)
SUB = RGBColor(0x81, 0x8C, 0xF8)
TEXT = RGBColor(0xE2, 0xE8, 0xF0)
ACCENT = RGBColor(0xF0, 0xAB, 0xFC)
MUTED = RGBColor(0x94, 0xA3, 0xB8)

W = Inches(13.333)
H = Inches(7.5)


def _bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG


def _textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    return tf


def _para(tf, text, size, color, bold=False, first=False, space=8):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.text = text
    p.space_after = Pt(space)
    for run in p.runs:
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.name = "Segoe UI"
    return p


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide)
    tf = _textbox(slide, Inches(1), Inches(2.2), Inches(11.3), Inches(3))
    _para(tf, "🐍 KnovaLand", 54, TITLE, bold=True, first=True)
    _para(tf, "Платформа самообучения с геймификацией", 26, TEXT, space=18)
    _para(tf, "ФИО, группа · Дипломный проект · 2026", 16, MUTED, space=14)
    _para(tf, "Django · DRF · PostgreSQL · JWT · Swagger · Docker", 14, ACCENT)


def content_slide(prs, title, lines):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide)
    tf = _textbox(slide, Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.0))
    _para(tf, title, 32, TITLE, bold=True, first=True)
    body = _textbox(slide, Inches(0.9), Inches(1.8), Inches(11.5), Inches(5.2))
    first = True
    for line in lines:
        if line.startswith("## "):
            _para(body, line[3:], 20, SUB, bold=True, first=first, space=6)
        elif line.startswith("> "):
            _para(body, line[2:], 16, ACCENT, first=first, space=6)
        elif line.startswith("- "):
            _para(body, "•  " + line[2:], 18, TEXT, first=first, space=6)
        elif line.startswith("* "):
            _para(body, "    ◦  " + line[2:], 16, MUTED, first=first, space=4)
        else:
            _para(body, line, 18, TEXT, first=first, space=6)
        first = False


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    title_slide(prs)

    content_slide(
        prs,
        "Проблема → Решение",
        [
            "## По ТЗ нужно:",
            "- Авторизация и аутентификация (JWT)",
            "- Управление контентом (admin + REST API)",
            "- Тестирование знаний (проверка отдельным запросом)",
            "- Роли и права (admin / teacher / student)",
            "## Что добавил от себя:",
            "> Обучение как игра — опыт, уровни, достижения, навыки, "
            "постепенное открытие и тренажёр.",
        ],
    )

    content_slide(
        prs,
        "Что реализовано",
        [
            "- 4 курса: Python · JavaScript · C# · Java — 28 уроков, 12 навыков",
            "- 3 роли: администратор · преподаватель · студент",
            "- Геймификация: опыт, уровни (общий и по курсу), достижения, рейтинг",
            "- Тренажёр с двумя типами вопросов: выбор и ручной ввод",
            "- «Мои скиллы» в профиле + полный список знаний",
            "- Постепенное открытие уроков",
        ],
    )

    content_slide(
        prs,
        "Архитектура",
        [
            "Модульный монолит на Django — единое ядро REST API + геймифицированный UI.",
            "- users — роли, JWT, регистрация, подтверждение email",
            "- courses — курсы → разделы",
            "- materials — уроки (материалы)",
            "- testing — тесты, вопросы, проверка ответов",
            "- gamification — опыт, уровни, навыки, покупки",
            "- web — UI на Django Templates + HTMX + Tailwind",
            "Клиент → Django/DRF → PostgreSQL",
        ],
    )

    content_slide(
        prs,
        "Стек и почему",
        [
            "- Django 5 + DRF — требование ТЗ, готовая админка и ORM",
            "- PostgreSQL — требование ТЗ, надёжное хранение",
            "- SimpleJWT — защита API токенами",
            "- drf-spectacular — автогенерируемый Swagger",
            "- django-filter — фильтрация и поиск",
            "- Docker / docker-compose — воспроизводимый запуск",
            "- UI: Django Templates + HTMX + Tailwind (без отдельного SPA)",
        ],
    )

    content_slide(
        prs,
        "Демонстрация — путь студента",
        [
            "- Главная → витрина из 4 курсов",
            "- Регистрация → письмо-подтверждение",
            "- Вход → запись на курс",
            "- Уроки открываются постепенно",
            "- Тренажёр: вопрос с выбором + ручным вводом",
            "- Результат → опыт, новый навык",
            "- Профиль → «Мои скиллы»   ·   Рейтинг игроков",
        ],
    )

    content_slide(
        prs,
        "Проверка знаний — отдельным запросом",
        [
            "POST /api/tests/{id}/submit/   (Authorization: Bearer <access>)",
            '{ "answers": [ {"question": 1, "choices": [3]}, {"question": 2, "text": "print"} ] }',
            '{ "score": 2, "total": 2, "percentage": 100.0, "is_passed": true }',
            "- Правильные ответы НИКОГДА не отдаются студенту",
            "- Два типа вопросов: выбор и ручной ввод",
            "- Расчёт баллов — на сервере",
        ],
    )

    content_slide(
        prs,
        "Роли и права",
        [
            "- Admin — полный доступ, управление пользователями",
            "- Teacher — CRUD только своих курсов, уроков и тестов",
            "- Student — просмотр материалов и прохождение тестов",
            "## Реализация:",
            "- Кастомные DRF-пермишены + поле роли",
            "- Покрыто автоматическими тестами",
        ],
    )

    content_slide(
        prs,
        "Результаты тестирования",
        [
            "## 95 тестов · 92% покрытия",
            "- Выше требуемых по ТЗ 75%",
            "- Покрыты: авторизация, права, CRUD, проверка тестов, геймификация",
            "- Соответствие PEP8 — flake8 = 0 нарушений",
            "- Проверено и на SQLite, и на PostgreSQL",
        ],
    )

    content_slide(
        prs,
        "Бизнес-ценность (STAR)",
        [
            "- S/T: ручная проверка тестов и низкая мотивация студентов",
            "- A: автопроверка + геймификация + монетизация курсов",
            "## R — цифры (модель на допущениях):",
            "- Автопроверка: −150 часов за поток",
            "- Экономия: ≈ 120 000 ₽ за поток",
            "- Вовлечённость: +50…60% completion rate",
            "- Монетизация: до 50 000 ₽/мес",
        ],
    )

    content_slide(
        prs,
        "Итоги и развитие · Спасибо!",
        [
            "## Готово:",
            "- Все требования ТЗ закрыты",
            "- 95 тестов, 92% покрытия, PEP8",
            "- Код в удалённом Git-репозитории",
            "- Docker-развёртывание",
            "## Дальше:",
            "- Аналитика для преподавателя · реальные платёжные шлюзы",
            "- Мобильное приложение и интеграции",
        ],
    )

    out = Path(__file__).resolve().parent.parent / "docs" / "KnovaLand_slides.pptx"
    prs.save(out)
    print(f"Saved: {out}")
    return out


if __name__ == "__main__":
    build()
